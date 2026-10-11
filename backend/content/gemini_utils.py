import os
import json
import logging
from django.conf import settings
import google.generativeai as genai

logger = logging.getLogger(__name__)

# Document Parsers
import pypdf
import docx
from pptx import Presentation

def extract_pdf_text(file_path):
    text = ""
    try:
        reader = pypdf.PdfReader(file_path)
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    except Exception as e:
        logger.exception(f"Error extracting PDF text from {file_path}")
    return text.strip()

def extract_pdf_content_and_diagrams(file_path, doc_id):
    """
    Extracts text page-by-page and detects/saves diagrams, flowcharts, block diagrams,
    tables, figures, network diagrams, graphs, charts, circuit diagrams, architecture diagrams,
    mathematical diagrams, and illustrations across multi-page PDFs.
    Returns: (extracted_text, diagrams_list)
    """
    import pymupdf
    from PIL import Image

    text_content = ""
    diagrams = []

    media_dir = os.path.join(settings.MEDIA_ROOT, 'document_diagrams', str(doc_id))
    os.makedirs(media_dir, exist_ok=True)
    media_url_base = f"{settings.MEDIA_URL}document_diagrams/{doc_id}/"

    try:
        doc = pymupdf.open(file_path)
        for page_idx in range(len(doc)):
            page_num = page_idx + 1
            page = doc[page_idx]

            page_text = page.get_text().strip()
            if page_text:
                text_content += f"\n--- Page {page_num} ---\n" + page_text + "\n"

            # 1. Detect embedded images
            image_list = page.get_images(full=True)
            saved_images_count = 0

            for img_idx, img_info in enumerate(image_list, start=1):
                try:
                    xref = img_info[0]
                    base_image = doc.extract_image(xref)
                    if not base_image:
                        continue

                    img_bytes = base_image["image"]
                    img_ext = base_image["ext"]
                    width = base_image.get("width", 0)
                    height = base_image.get("height", 0)

                    # Filter out tiny icon decorations (less than 60x60 pixels)
                    if width >= 60 and height >= 60:
                        img_filename = f"page_{page_num}_diag_{img_idx}.{img_ext}"
                        abs_img_path = os.path.join(media_dir, img_filename)

                        with open(abs_img_path, "wb") as f:
                            f.write(img_bytes)

                        rel_media_path = f"document_diagrams/{doc_id}/{img_filename}"
                        img_url = f"{media_url_base}{img_filename}"

                        diagrams.append({
                            "id": f"diag-p{page_num}-{img_idx}",
                            "page_number": page_num,
                            "image_url": img_url,
                            "image_path": abs_img_path,
                            "relative_path": rel_media_path,
                            "caption": f"Page {page_num} Figure / Diagram {img_idx}",
                            "page_context": page_text[:500]
                        })
                        saved_images_count += 1
                        if saved_images_count >= 5:
                            break
                except Exception as img_err:
                    logger.warning(f"Error extracting image {img_idx} from page {page_num}: {img_err}")

            # 2. Check for vector drawings (flowcharts, tables, graphs, charts, block diagrams)
            drawings = page.get_drawings()
            has_visual_keywords = any(kw in page_text.lower() for kw in [
                'figure', 'fig', 'diagram', 'flowchart', 'block diagram', 'table', 'chart',
                'graph', 'architecture', 'circuit', 'illustration', 'model'
            ])

            if len(drawings) >= 10 or (has_visual_keywords and saved_images_count == 0):
                try:
                    pix = page.get_pixmap(dpi=150)
                    img_filename = f"page_{page_num}_render.png"
                    abs_img_path = os.path.join(media_dir, img_filename)
                    pix.save(abs_img_path)

                    rel_media_path = f"document_diagrams/{doc_id}/{img_filename}"
                    img_url = f"{media_url_base}{img_filename}"

                    diagrams.append({
                        "id": f"diag-p{page_num}-render",
                        "page_number": page_num,
                        "image_url": img_url,
                        "image_path": abs_img_path,
                        "relative_path": rel_media_path,
                        "caption": f"Page {page_num} Visual Page / Diagram",
                        "page_context": page_text[:500]
                    })
                except Exception as render_err:
                    logger.warning(f"Error rendering page {page_num} pixmap: {render_err}")

        doc.close()

    except Exception as e:
        logger.exception(f"Error extracting PDF content and diagrams from {file_path}")
        if not text_content:
            text_content = extract_pdf_text(file_path)

    return text_content.strip(), diagrams

def extract_docx_text(file_path):
    text = ""
    try:
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
    except Exception as e:
        logger.exception(f"Error extracting DOCX text from {file_path}")
    return text.strip()

def extract_pptx_text(file_path):
    text = ""
    try:
        prs = Presentation(file_path)
        for idx, slide in enumerate(prs.slides):
            text += f"\n--- Slide {idx+1} ---\n"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"
    except Exception as e:
        logger.exception(f"Error extracting PPTX text from {file_path}")
    return text.strip()

def generate_summary_and_keywords_from_text(text):
    """
    Sends raw extracted document text to Gemini to generate 
    a detailed summary and keyword tags.
    Returns: (summary_text, keywords_list)
    """
    configure_gemini()
    try:
        model = genai.GenerativeModel("gemini-3.5-flash-lite")
        
        prompt = (
            "You are an expert academic scribe and AI learning assistant. "
            "Analyze the following lecture or course document text. "
            "1. Generate a structured, comprehensive summary of the core concepts discussed. "
            "2. Extract a list of 5-10 key terminology/topic tags (keywords). "
            "Respond ONLY with a JSON object in this format: "
            "{\n"
            "  \"summary\": \"string (comprehensive summary with markdown bullet points)\",\n"
            "  \"keywords\": [\"string\", \"string\"]\n"
            "}\n\n"
            f"--- TEXT CONTENT ---\n{text}"
        )
        
        response = model.generate_content(
            prompt,
            generation_config={"response_mime_type": "application/json"}
        )
        result = json.loads(response.text)
        return (
            result.get("summary", ""),
            result.get("keywords", [])
        )
    except Exception as e:
        logger.exception("Error generating summary from text via Gemini")
        return (
            f"Failed to generate summary. Details: {str(e)}",
            ["Error"]
        )


def configure_gemini():
    """Initializes the Gemini API client with settings API key."""
    api_key = getattr(settings, 'GEMINI_API_KEY', '')
    if not api_key:
        logger.warning("GEMINI_API_KEY not configured in Django settings.")
    genai.configure(api_key=api_key)

def transcribe_and_summarize_audio(file_path):
    """
    Uploads audio file to Gemini API, transcribes it, 
    and generates a detailed summary and keyword tags.
    Returns: (transcript_text, summary_text, keywords_list)
    """
    configure_gemini()
    
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Audio file not found at path: {file_path}")

    try:
        # Upload file to Gemini File API
        logger.info(f"Uploading file to Gemini: {file_path}")
        audio_file = genai.upload_file(path=file_path)
        logger.info(f"Uploaded successfully. File URI: {audio_file.uri}")

        # Choose the flash model for speed and cost-effectiveness
        model = genai.GenerativeModel("gemini-3.5-flash-lite")

        prompt = (
            "You are an expert academic scribe and AI learning assistant. "
            "Analyze the attached audio lecture file. "
            "1. Transcribe the audio lecture fully and word-for-word. "
            "2. Generate a structured, comprehensive summary of the core concepts discussed. "
            "3. Extract a list of 5-10 key terminology/topic tags (keywords). "
            "Respond ONLY with a JSON object in this format: "
            "{\n"
            "  \"transcript\": \"string (entire transcription)\",\n"
            "  \"summary\": \"string (comprehensive summary with markdown bullet points)\",\n"
            "  \"keywords\": [\"string\", \"string\"]\n"
            "}"
        )

        response = model.generate_content(
            [audio_file, prompt],
            generation_config={"response_mime_type": "application/json"}
        )

        # Cleanup the file from Gemini cloud storage
        try:
            genai.delete_file(audio_file.name)
        except Exception as cleanup_err:
            logger.warning(f"Failed to delete Gemini temporary file: {cleanup_err}")

        # Parse output
        result = json.loads(response.text)
        return (
            result.get("transcript", ""),
            result.get("summary", ""),
            result.get("keywords", [])
        )

    except Exception as e:
        logger.exception("Error processing audio via Gemini API")
        # Return fallback text to prevent application crashes
        return (
            "[Transcription failed due to API error or invalid key]",
            f"Failed to generate summary. Details: {str(e)}",
            ["Error"]
        )

def generate_study_materials(transcript):
    """
    Generates educational flashcards and interactive quiz questions 
    based on a lecture transcript.
    Returns: (flashcards_list, quiz_questions_list)
    """
    configure_gemini()
    
    if not transcript or len(transcript.strip()) < 20:
        return [], []

    try:
        model = genai.GenerativeModel("gemini-3.5-flash-lite")

        prompt = (
            "Analyze the following lecture transcript and create study materials:\n"
            f"--- TRANSCRIPT ---\n{transcript}\n-------------------\n\n"
            "Create:\n"
            "1. A set of 5-8 flashcards covering key terms and definitions. Each flashcard has a front and a back.\n"
            "2. A quiz of 5 multiple-choice questions (MCQs) and 2 true/false questions. "
            "Each question should have a question_text, a question_type ('MCQ' or 'TRUE_FALSE'), "
            "a list of options (for MCQs only; use empty list for true/false), and correct_answer (string matching options or 'True'/'False').\n\n"
            "You MUST respond ONLY with a JSON object matching this schema:\n"
            "{\n"
            "  \"flashcards\": [\n"
            "     { \"front\": \"question or term\", \"back\": \"answer or definition\" }\n"
            "  ],\n"
            "  \"quiz_questions\": [\n"
            "     {\n"
            "       \"question_text\": \"question content\",\n"
            "       \"question_type\": \"MCQ\",\n"
            "       \"options\": [\"Option A\", \"Option B\", \"Option C\", \"Option D\"],\n"
            "       \"correct_answer\": \"Option B\"\n"
            "     }\n"
            "  ]\n"
            "}"
        )

        response = model.generate_content(
            prompt,
            generation_config={"response_mime_type": "application/json"}
        )

        result = json.loads(response.text)
        return (
            result.get("flashcards", []),
            result.get("quiz_questions", [])
        )

    except Exception as e:
        logger.exception("Error generating study aids via Gemini API")
        return [], []


def transcribe_audio_message(file_path):
    """
    Transcribes a short audio message file using Gemini API.
    """
    configure_gemini()
    if not os.path.exists(file_path):
        return ""
    try:
        audio_file = genai.upload_file(path=file_path)
        model = genai.GenerativeModel("gemini-3.5-flash-lite")
        response = model.generate_content(
            [audio_file, "Transcribe this audio message word for word. Do not add comments, labels, or explanations. Respond with just the transcription."]
        )
        try:
            genai.delete_file(audio_file.name)
        except Exception:
            pass
        return response.text.strip()
    except Exception as e:
        logger.exception("Error transcribing voice message chunk")
        return "[Audio transcription failed]"


TELUGU_PROMPT_INSTRUCTION = (
    "You are explaining technical topics to an engineering student in Telugu.\n\n"
    "Use simple, natural, conversational Telugu (వాడుక భాష).\n"
    "Explain as a friendly college lecturer would explain in class ('మన lecturer class లో easy గా explain చేస్తున్నట్టు').\n\n"
    "Avoid grandhika Telugu, literary Telugu, Sanskrit-heavy vocabulary, and overly formal bookish sentences.\n\n"
    "Use common everyday Telugu words.\n"
    "Keep technical terms such as Computer, Database, Network, Algorithm, API, Python, Machine Learning, Server, Table, CPU, etc. in English script when that makes the explanation easier and more natural.\n\n"
    "Do not translate English sentences word-by-word.\n"
    "Explain the concept meaning naturally in Telugu.\n\n"
    "Use short sentences and simple real-world examples.\n"
    "The student should understand the concept immediately after listening to the explanation."
)


def retrieve_relevant_pdf_chunks(full_text, query, top_k=4, chunk_size=800, overlap=150):
    """
    RAG Chunk Retrieval Engine:
    Splits uploaded PDF extracted text into overlapping chunks, scores them by keyword relevance
    to the student's doubt query, and returns the most relevant retrieved context.
    """
    if not full_text or len(full_text.strip()) == 0:
        return "", False

    import re
    stop_words = {'what', 'is', 'the', 'a', 'an', 'in', 'on', 'at', 'for', 'to', 'of', 'and', 'or', 'how', 'why', 'can', 'you', 'explain', 'tell', 'me', 'about', 'i', 'dont', 'understand', 'meaning'}
    query_words = [w.lower() for w in re.findall(r'\w+', query) if w.lower() not in stop_words and len(w) > 1]
    if not query_words:
        query_words = [w.lower() for w in re.findall(r'\w+', query) if len(w) > 1]

    text_len = len(full_text)
    chunks = []
    start = 0
    while start < text_len:
        end = min(start + chunk_size, text_len)
        chunks.append(full_text[start:end])
        if end == text_len:
            break
        start += (chunk_size - overlap)

    if not chunks:
        return full_text[:3000], True

    scored_chunks = []
    for idx, chunk in enumerate(chunks):
        chunk_lower = chunk.lower()
        score = 0
        for qw in query_words:
            score += chunk_lower.count(qw) * 2
            if query.lower() in chunk_lower:
                score += 10
        scored_chunks.append((score, idx, chunk))

    scored_chunks.sort(key=lambda x: x[0], reverse=True)
    top_score = scored_chunks[0][0]

    if top_score == 0:
        # No keyword match found in document chunks
        return full_text[:3000], False

    selected_scored = scored_chunks[:top_k]
    selected_scored.sort(key=lambda x: x[1])
    retrieved_text = "\n\n".join([c for _, _, c in selected_scored])
    return retrieved_text, True


def explain_text_segment(text, language='en'):
    """
    Explains the given text segment in English or simple spoken Telugu using Gemini.
    """
    configure_gemini()
    try:
        model = genai.GenerativeModel("gemini-3.5-flash-lite")
        lang_str = str(language).lower().strip()
        is_telugu = lang_str in ['telugu', 'te', 'te-in', 'te_in'] or lang_str.startswith('te')
        
        if is_telugu:
            prompt = (
                f"{TELUGU_PROMPT_INSTRUCTION}\n"
                f"Explain the following academic engineering text segment in simple conversational Telugu:\n\n"
                f"--- TEXT SEGMENT ---\n{text}"
            )
        else:
            prompt = (
                f"You are an encouraging academic teacher. Explain the following text segment "
                f"clearly and simply in English. Output ONLY the explanation. Do not add system logs or headers.\n\n"
                f"--- TEXT SEGMENT ---\n{text}"
            )
        
        response = model.generate_content(prompt)
        return response.text.strip()
    except Exception as e:
        logger.exception(f"Error explaining segment in {language}")
        return f"Error generating explanation: {str(e)}"


def chat_with_document_context(question, doc_context, history=None):
    """
    Answers a student's question based on the document text context and conversation history.
    """
    configure_gemini()
    try:
        model = genai.GenerativeModel("gemini-3.5-flash-lite")
        
        history_str = ""
        if history:
            for chat in history:
                role = "Student" if chat.get('role') == 'user' else "Assistant"
                history_str += f"{role}: {chat.get('text')}\n"

        prompt = (
            "You are an intelligent educational chat companion. Answer the student's question "
            "based strictly on the document context provided below. If the answer cannot be "
            "found in the context, use your general knowledge but mention it is supplementary information.\n\n"
            f"--- DOCUMENT CONTEXT ---\n{doc_context}\n\n"
            f"--- CONVERSATION HISTORY ---\n{history_str}"
            f"Student: {question}\n"
            "Assistant:"
        )

        response = model.generate_content(prompt)
        return response.text.strip()
    except Exception as e:
        logger.exception("Error chatting with document context")
        return f"Apologies, I encountered an issue parsing that query: {str(e)}"


def call_gemini_multimodal(prompt_prefix, full_text, diagrams=None, prompt_suffix="", max_chunk_size=90000):
    """
    Submits text and diagram images to Gemini for multimodal visual understanding & analysis.
    Supports up to 10 visual diagram images per request.
    """
    from PIL import Image
    configure_gemini()
    model = genai.GenerativeModel("gemini-3.5-flash-lite")

    image_inputs = []
    if diagrams and isinstance(diagrams, list):
        for diag in diagrams[:10]:
            img_path = diag.get('image_path')
            if img_path and os.path.exists(img_path):
                try:
                    pil_img = Image.open(img_path)
                    image_inputs.append(pil_img)
                except Exception as img_err:
                    logger.warning(f"Could not load image from {img_path}: {img_err}")

    diagram_instruction = (
        "\nIMPORTANT MANDATORY REQUIREMENTS FOR DIAGRAMS & VISUAL CONTENT:\n"
        "1. The document contains visual diagrams, figures, flowcharts, block diagrams, tables, graphs, charts, or illustrations.\n"
        "2. You MUST include a dedicated section titled '### 📊 Diagrams & Visuals' in your response.\n"
        "3. For EVERY important diagram/figure found, explain:\n"
        "   - Diagram/Figure name & page number\n"
        "   - What the diagram represents\n"
        "   - Main components shown and what each component does\n"
        "   - How the components are connected / flow or process shown\n"
        "   - Important labels, values, or numbers\n"
        "   - Why the diagram is important to the overall topic\n"
        "4. Do NOT invent or hallucinate information that is not visible in the visual diagram.\n"
        "5. If a diagram or text in a visual is blurry or cannot be clearly read, explicitly state:\n"
        "   'This diagram is present in the PDF, but some parts are not clear enough to interpret accurately.'\n"
        "6. Connect the text explanation with the visual explanation where relevant.\n"
    )

    full_prompt = f"{prompt_prefix}\n{diagram_instruction}\nDocument Text Context:\n{full_text[:max_chunk_size]}\n\n{prompt_suffix}"

    try:
        if image_inputs:
            contents = image_inputs + [full_prompt]
            response = model.generate_content(contents)
        else:
            response = model.generate_content(full_prompt)

        return response.text.strip()
    except Exception as e:
        logger.exception("Multimodal call to Gemini failed, retrying with text only...")
        try:
            response = model.generate_content(full_prompt)
            return response.text.strip()
        except Exception as text_err:
            return f"Gemini API error: {str(text_err)}"


def answer_doubt_with_rag(question, document_text=None, history=None, language='english', diagrams=None):
    """
    Answers a student's doubt using document RAG chunk retrieval, Gemini generation,
    and simple spoken Telugu/English prompts. Includes diagram visual context where relevant.
    """
    from PIL import Image
    configure_gemini()
    try:
        model = genai.GenerativeModel("gemini-3.5-flash-lite")
        
        retrieved_context = ""
        context_found = True
        
        if document_text and len(document_text.strip()) > 0:
            retrieved_context, context_found = retrieve_relevant_pdf_chunks(document_text, question)

        history_str = ""
        if history:
            for item in history:
                role = "Student" if item.get('sender') in ['USER', 'STUDENT', 'user'] else "Lecturer"
                history_str += f"{role}: {item.get('text_content') or item.get('text')}\n"

        lang_str = str(language).lower().strip()
        is_telugu = lang_str in ['telugu', 'te', 'te-in', 'te_in'] or lang_str.startswith('te')

        # Collect relevant diagram images if available
        image_inputs = []
        retrieved_diagrams = []
        if diagrams and isinstance(diagrams, list):
            for diag in diagrams[:3]:
                img_path = diag.get('image_path')
                if img_path and os.path.exists(img_path):
                    try:
                        pil_img = Image.open(img_path)
                        image_inputs.append(pil_img)
                        retrieved_diagrams.append(diag)
                    except Exception as err:
                        logger.warning(f"Error loading RAG diagram image: {err}")

        diagram_note = ""
        if image_inputs:
            diagram_note = (
                f"\n[Visual Context: {len(image_inputs)} diagram(s)/figure(s) from the document are attached to this prompt. "
                "Examine the diagram visual structure (flowchart, block diagram, table, graph, circuit, etc.), "
                "explain it simply and clearly, and connect it to the doubt explanation.]\n"
            )

        if is_telugu:
            doc_instruction = ""
            if document_text:
                if context_found:
                    doc_instruction = (
                        f"--- UPLOADED STUDY MATERIAL CONTEXT ---\n{retrieved_context}\n\n"
                        "Instructions: Base your answer primarily on the uploaded study material context above. "
                        "Explain in simple conversational Telugu + English technical terms.\n"
                    )
                else:
                    doc_instruction = (
                        f"--- UPLOADED STUDY MATERIAL CONTEXT ---\n{retrieved_context[:2000]}\n\n"
                        "Instructions: The answer for the student's doubt is NOT specifically found in the uploaded study material. "
                        "You MUST start your response by explicitly stating in simple Telugu:\n"
                        "'ఈ సమాచారం మీరు upload చేసిన document లో స్పష్టంగా లేదు, కానీ సాధారణంగా దీని అర్థం...'\n"
                        "and then provide a simple general explanation with an example.\n"
                    )

            prompt = (
                f"{TELUGU_PROMPT_INSTRUCTION}\n"
                f"{diagram_note}\n"
                f"{doc_instruction}\n"
                f"--- CONVERSATION HISTORY ---\n{history_str}\n"
                f"Student Doubt: {question}\n"
                f"Lecturer Answer in Simple Telugu:"
            )
        else:
            doc_instruction = ""
            if document_text:
                if context_found:
                    doc_instruction = (
                        f"--- UPLOADED STUDY MATERIAL CONTEXT ---\n{retrieved_context}\n\n"
                        "Instructions: Base your answer primarily on the uploaded study material context above.\n"
                    )
                else:
                    doc_instruction = (
                        f"--- UPLOADED STUDY MATERIAL CONTEXT ---\n{retrieved_context[:2000]}\n\n"
                        "Instructions: The answer for the student's doubt is NOT found in the uploaded study material. "
                        "You MUST start your response by clearly stating:\n"
                        "'This specific detail is not found in your uploaded study material, but generally speaking...'\n"
                        "and then provide a clear, simple general explanation with an example.\n"
                    )

            prompt = (
                "You are an encouraging academic engineering lecturer clearing a student's doubt.\n"
                "Rules:\n"
                "- Explain clearly, simply, and step-by-step with practical real-world examples.\n"
                "- If the student asks a follow-up question or says 'I don't understand', explain it in an even simpler, intuitive way.\n"
                f"{diagram_note}\n"
                f"{doc_instruction}\n"
                f"--- CONVERSATION HISTORY ---\n{history_str}\n"
                f"Student Doubt: {question}\n"
                f"Lecturer Answer:"
            )

        if image_inputs:
            contents = image_inputs + [prompt]
            response = model.generate_content(contents)
        else:
            response = model.generate_content(prompt)

        return response.text.strip(), context_found, retrieved_context, retrieved_diagrams

    except Exception as e:
        logger.exception("Error in answer_doubt_with_rag")
        err_msg = (
            f"క్షమించండి, doubt answer generate చేయడంలో సమస్య వచ్చింది: {str(e)}"
            if language == 'telugu' else
            f"Apologies, I encountered an issue processing your doubt: {str(e)}"
        )
        return err_msg, False, "", []



